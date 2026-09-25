"""
Éditeur de documents existants

Applique des modifications à un DOCX, XLSX, PPTX, PDF ou fichier texte joint
par l'utilisateur. **Le fichier source n'est jamais modifié** : le résultat
est toujours écrit dans ``outputs/documents/<nom>_modifie.<ext>``.

Le PDF fait exception au niveau de la technique : il n'est pas éditable en
place, donc son texte est extrait, modifié, puis un nouveau PDF est régénéré
par ``DocumentGenerator`` — la mise en page d'origine est perdue, et le
message de retour le dit explicitement.
"""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from generators.document_generator import DOCUMENTS_DIR, DocumentGenerator
from generators.markdown_document import parse_markdown

# Suffixe ajouté au nom du fichier d'origine pour la copie modifiée.
_MODIFIED_SUFFIX = "_modifie"

# Opérations acceptées, par famille de format.
_TEXTUAL_ACTIONS = {"replace_text", "append_markdown", "replace_section", "delete_paragraph"}
_SHEET_ACTIONS = {"set_cell", "append_row"}
_SLIDE_ACTIONS = {"append_slide"}

SUPPORTED_EXTENSIONS = {".docx", ".xlsx", ".xlsm", ".pptx", ".pdf",
                        ".md", ".markdown", ".txt", ".csv"}


class DocumentEditError(ValueError):
    """Opération d'édition impossible ou mal formée."""


class DocumentEditor:
    """
    Applique des opérations d'édition à un document, sur une copie.
    """

    def __init__(self, generator: Optional[DocumentGenerator] = None):
        """
        Args:
            generator: Générateur réutilisé pour régénérer un PDF modifié.
                Instancié à la demande si absent.
        """
        self._generator = generator

    # ── API publique ───────────────────────────────────────────────────────

    def edit(
        self,
        path: str,
        operations: List[Dict[str, Any]],
        output_name: str = "",
    ) -> Dict[str, Any]:
        """
        Applique une liste d'opérations et écrit une copie modifiée.

        Args:
            path: Chemin du document à modifier (jamais écrasé)
            operations: Liste de dicts `{"action": …, …}`
            output_name: Nom du fichier de sortie (sans dossier). Dérivé du
                nom d'origine s'il est vide.

        Returns:
            {success, file_path, file_name, source_path, applied, notes}
            ou {success: False, error}
        """
        source = Path(path).expanduser()
        if not source.exists():
            return {"success": False, "error": f"Fichier introuvable : {path}"}

        suffix = source.suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            return {
                "success": False,
                "error": (
                    f"Format non modifiable : {suffix}. "
                    f"Formats acceptés : {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
                ),
            }

        operations = _normalize_operations(operations)
        if not operations:
            return {"success": False, "error": "Aucune opération d'édition fournie."}

        destination = self._destination(source, output_name)
        destination.parent.mkdir(parents=True, exist_ok=True)

        try:
            if suffix == ".docx":
                applied, notes = self._edit_docx(source, destination, operations)
            elif suffix in (".xlsx", ".xlsm"):
                applied, notes = self._edit_xlsx(source, destination, operations)
            elif suffix == ".pptx":
                applied, notes = self._edit_pptx(source, destination, operations)
            elif suffix == ".pdf":
                applied, notes = self._edit_pdf(source, destination, operations)
            else:
                applied, notes = self._edit_text(source, destination, operations)
        except DocumentEditError as exc:
            return {"success": False, "error": str(exc)}
        except (OSError, ValueError, KeyError, ImportError) as exc:
            return {"success": False, "error": f"Erreur lors de la modification : {exc}"}

        if applied == 0:
            return {
                "success": False,
                "error": (
                    "Aucune opération n'a pu être appliquée : le texte recherché "
                    "est introuvable dans le document."
                ),
            }

        return {
            "success": True,
            "file_path": str(destination),
            "file_name": destination.name,
            "source_path": str(source),
            "applied": applied,
            "notes": notes,
        }

    # ── DOCX ───────────────────────────────────────────────────────────────

    def _edit_docx(
        self, source: Path, destination: Path, operations: List[Dict[str, Any]]
    ) -> Tuple[int, List[str]]:
        """Modifie un DOCX en conservant sa mise en forme."""
        import docx

        document = docx.Document(str(source))
        applied = 0
        notes: List[str] = []

        for operation in operations:
            action = operation["action"]

            if action == "replace_text":
                find, replace = _require_find_replace(operation)
                applied += _docx_replace_text(document, find, replace)

            elif action == "delete_paragraph":
                needle = _require(operation, "contains")
                applied += _docx_delete_paragraphs(document, needle)

            elif action == "append_markdown":
                _docx_append_markdown(document, _require(operation, "content"))
                applied += 1

            elif action == "replace_section":
                heading = _require(operation, "heading")
                count = _docx_replace_section(
                    document, heading, operation.get("content", "")
                )
                if count:
                    applied += count
                else:
                    notes.append(f"Section « {heading} » introuvable.")

            else:
                notes.append(f"Opération « {action} » ignorée pour un DOCX.")

        document.save(str(destination))
        return applied, notes

    # ── XLSX ───────────────────────────────────────────────────────────────

    def _edit_xlsx(
        self, source: Path, destination: Path, operations: List[Dict[str, Any]]
    ) -> Tuple[int, List[str]]:
        """Modifie un classeur Excel cellule par cellule."""
        from openpyxl import load_workbook

        workbook = load_workbook(str(source))
        applied = 0
        notes: List[str] = []

        for operation in operations:
            action = operation["action"]

            if action == "set_cell":
                sheet = _resolve_sheet(workbook, operation.get("sheet", ""))
                reference = _require(operation, "cell")
                sheet[reference] = operation.get("value", "")
                applied += 1

            elif action == "append_row":
                sheet = _resolve_sheet(workbook, operation.get("sheet", ""))
                values = operation.get("values")
                if not isinstance(values, (list, tuple)):
                    raise DocumentEditError(
                        "append_row attend une liste 'values' de valeurs de cellules."
                    )
                sheet.append(list(values))
                applied += 1

            elif action == "replace_text":
                find, replace = _require_find_replace(operation)
                for sheet in workbook.worksheets:
                    for row in sheet.iter_rows():
                        for cell in row:
                            if isinstance(cell.value, str) and find in cell.value:
                                cell.value = cell.value.replace(find, replace)
                                applied += 1

            else:
                notes.append(f"Opération « {action} » ignorée pour un classeur.")

        workbook.save(str(destination))
        return applied, notes

    # ── PPTX ───────────────────────────────────────────────────────────────

    def _edit_pptx(
        self, source: Path, destination: Path, operations: List[Dict[str, Any]]
    ) -> Tuple[int, List[str]]:
        """Modifie une présentation : texte des formes et ajout de diapos."""
        from pptx import Presentation

        presentation = Presentation(str(source))
        applied = 0
        notes: List[str] = []

        for operation in operations:
            action = operation["action"]

            if action == "replace_text":
                find, replace = _require_find_replace(operation)
                applied += _pptx_replace_text(presentation, find, replace)

            elif action == "append_slide":
                _pptx_append_slide(
                    presentation,
                    operation.get("title", ""),
                    operation.get("content", ""),
                )
                applied += 1

            elif action == "delete_paragraph":
                needle = _require(operation, "contains")
                applied += _pptx_delete_paragraphs(presentation, needle)

            else:
                notes.append(f"Opération « {action} » ignorée pour une présentation.")

        presentation.save(str(destination))
        return applied, notes

    # ── PDF ────────────────────────────────────────────────────────────────

    def _edit_pdf(
        self, source: Path, destination: Path, operations: List[Dict[str, Any]]
    ) -> Tuple[int, List[str]]:
        """
        Régénère un PDF modifié à partir de son texte extrait.

        Un PDF n'est pas éditable en place : la mise en page d'origine (polices,
        colonnes, images) est perdue. C'est signalé dans les notes.
        """
        from processors.pdf_processor import PDFProcessor

        text = PDFProcessor().extract_text(str(source))
        if not text or not text.strip():
            raise DocumentEditError(
                "Texte du PDF illisible (document scanné ou protégé) : "
                "modification impossible."
            )

        text, applied = _apply_text_operations(text, operations)

        # llm=False : la régénération ne fait que du rendu, inutile de sonder Ollama.
        generator = self._generator or DocumentGenerator(llm=False)
        generator.build(parse_markdown(text), "pdf", source.stem, destination)

        return applied, [
            "Un PDF n'est pas modifiable en place : un nouveau PDF a été "
            "régénéré à partir du texte extrait, sans la mise en page d'origine."
        ]

    # ── Texte / Markdown / CSV ─────────────────────────────────────────────

    def _edit_text(
        self, source: Path, destination: Path, operations: List[Dict[str, Any]]
    ) -> Tuple[int, List[str]]:
        """Modifie un fichier texte, Markdown ou CSV."""
        try:
            text = source.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = source.read_text(encoding="latin-1")

        text, applied = _apply_text_operations(text, operations)
        destination.write_text(text, encoding="utf-8")
        return applied, []

    # ── Chemins ────────────────────────────────────────────────────────────

    @staticmethod
    def _destination(source: Path, output_name: str) -> Path:
        """
        Construit le chemin de la copie modifiée dans outputs/documents/.

        Le nom est suffixé d'un compteur si une copie existe déjà, pour ne
        jamais écraser une modification précédente.
        """
        if output_name:
            stem = Path(output_name).stem
            suffix = Path(output_name).suffix or source.suffix
        else:
            stem = f"{source.stem}{_MODIFIED_SUFFIX}"
            suffix = source.suffix

        candidate = DOCUMENTS_DIR / f"{stem}{suffix}"
        counter = 2
        while candidate.exists():
            candidate = DOCUMENTS_DIR / f"{stem}_{counter}{suffix}"
            counter += 1
        return candidate


# ══════════════════════════════════════════════════════════════════════════
#  Normalisation des opérations
# ══════════════════════════════════════════════════════════════════════════

# Synonymes tolérés : les modèles locaux nomment les champs de façons variées.
_ACTION_ALIASES = {
    "replace": "replace_text", "remplacer": "replace_text",
    "find_replace": "replace_text", "substitute": "replace_text",
    "append": "append_markdown", "add_text": "append_markdown",
    "ajouter": "append_markdown", "add_section": "append_markdown",
    "delete": "delete_paragraph", "remove": "delete_paragraph",
    "supprimer": "delete_paragraph",
    "update_section": "replace_section", "edit_section": "replace_section",
    "add_slide": "append_slide", "new_slide": "append_slide",
    "set_value": "set_cell", "write_cell": "set_cell",
    "add_row": "append_row",
}

_FIELD_ALIASES = {
    "search": "find", "old": "find", "old_text": "find", "from": "find",
    "target": "find", "chercher": "find", "texte": "find",
    "new": "replace", "new_text": "replace", "to": "replace",
    "replacement": "replace", "remplacement": "replace",
    "text": "content", "body": "content", "contenu": "content",
    "markdown": "content", "titre": "title", "section": "heading",
    "feuille": "sheet", "cellule": "cell", "valeur": "value",
}


def _normalize_operations(operations: Any) -> List[Dict[str, Any]]:
    """
    Normalise la liste d'opérations reçue d'un appel d'outil.

    Accepte un dict unique, une liste de dicts, des noms d'action et de champs
    alternatifs. Les entrées inexploitables sont écartées silencieusement.
    """
    if isinstance(operations, dict):
        operations = [operations]
    if not isinstance(operations, (list, tuple)):
        return []

    normalized: List[Dict[str, Any]] = []
    for raw in operations:
        if not isinstance(raw, dict):
            continue
        entry = {_FIELD_ALIASES.get(key, key): value for key, value in raw.items()}
        action = str(entry.get("action", "")).strip().lower()
        action = _ACTION_ALIASES.get(action, action)
        if action not in (_TEXTUAL_ACTIONS | _SHEET_ACTIONS | _SLIDE_ACTIONS):
            continue
        entry["action"] = action
        normalized.append(entry)
    return normalized


def _require(operation: Dict[str, Any], field: str) -> str:
    """Récupère un champ obligatoire, en signalant clairement son absence."""
    value = operation.get(field)
    if value is None or str(value) == "":
        raise DocumentEditError(
            f"Champ « {field} » manquant pour l'opération « {operation['action']} »."
        )
    return str(value)


def _require_find_replace(operation: Dict[str, Any]) -> Tuple[str, str]:
    """Récupère le couple (find, replace) d'un remplacement de texte."""
    return _require(operation, "find"), str(operation.get("replace", ""))


def _apply_text_operations(text: str, operations: List[Dict[str, Any]]) -> Tuple[str, int]:
    """Applique les opérations textuelles à une chaîne Markdown/texte."""
    applied = 0

    for operation in operations:
        action = operation["action"]

        if action == "replace_text":
            find, replace = _require_find_replace(operation)
            occurrences = text.count(find)
            if occurrences:
                text = text.replace(find, replace)
                applied += occurrences

        elif action == "append_markdown":
            text = text.rstrip() + "\n\n" + _require(operation, "content").strip() + "\n"
            applied += 1

        elif action == "delete_paragraph":
            needle = _require(operation, "contains")
            kept = [line for line in text.splitlines() if needle not in line]
            removed = len(text.splitlines()) - len(kept)
            if removed:
                text = "\n".join(kept)
                applied += removed

        elif action == "replace_section":
            heading = _require(operation, "heading")
            text, count = _replace_text_section(
                text, heading, operation.get("content", "")
            )
            applied += count

    return text, applied


def _replace_text_section(text: str, heading: str, content: str) -> Tuple[str, int]:
    """
    Remplace le corps d'une section Markdown, de son titre au titre suivant.

    Le titre lui-même est conservé.
    """
    lines = text.splitlines()
    start = None
    level = 0

    for index, line in enumerate(lines):
        stripped = line.strip()
        if stripped.startswith("#") and heading.lower() in stripped.lower():
            start = index
            level = len(stripped) - len(stripped.lstrip("#"))
            break

    if start is None:
        return text, 0

    end = len(lines)
    for index in range(start + 1, len(lines)):
        stripped = lines[index].strip()
        if stripped.startswith("#"):
            next_level = len(stripped) - len(stripped.lstrip("#"))
            if next_level <= level:
                end = index
                break

    replacement = [lines[start], ""] + content.strip().splitlines() + [""]
    return "\n".join(lines[:start] + replacement + lines[end:]), 1


# ══════════════════════════════════════════════════════════════════════════
#  Aides DOCX
# ══════════════════════════════════════════════════════════════════════════


def _docx_iter_paragraphs(document):
    """Itère sur tous les paragraphes, corps et cellules de tableau compris."""
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _replace_in_paragraph(paragraph, find: str, replace: str) -> int:
    """
    Remplace un texte dans un paragraphe en préservant au mieux le formatage.

    Deux passes, parce que Word et PowerPoint découpent régulièrement une
    phrase sur plusieurs runs (correcteur orthographique, suivi de
    modifications) :

    1. **run par run** — cas courant, où l'occurrence tient dans un seul run :
       la mise en forme de ce run (gras, couleur…) est intégralement conservée ;
    2. **repli au niveau du paragraphe** — seulement si l'occurrence est
       découpée entre plusieurs runs : le texte est réinjecté dans le premier
       run et les suivants sont vidés. Le paragraphe prend alors le formatage
       de son premier run, ce qu'aucune passe run par run ne peut éviter.
    """
    count = 0
    for run in paragraph.runs:
        if find in run.text:
            count += run.text.count(find)
            run.text = run.text.replace(find, replace)

    full_text = "".join(run.text for run in paragraph.runs) or paragraph.text
    if find not in full_text:
        return count

    # L'occurrence chevauche plusieurs runs : repli paragraphe.
    count += full_text.count(find)
    updated = full_text.replace(find, replace)
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(updated)
    return count


def _docx_replace_text(document, find: str, replace: str) -> int:
    """Remplace un texte dans tout le document, tableaux compris."""
    return sum(
        _replace_in_paragraph(paragraph, find, replace)
        for paragraph in _docx_iter_paragraphs(document)
    )


def _docx_delete_paragraphs(document, needle: str) -> int:
    """Supprime les paragraphes contenant un texte donné."""
    count = 0
    for paragraph in list(document.paragraphs):
        if needle in paragraph.text:
            element = paragraph._p  # noqa: SLF001 - pas d'API publique de suppression
            element.getparent().remove(element)
            count += 1
    return count


def _docx_append_markdown(document, content: str) -> None:
    """Ajoute du contenu Markdown rendu à la fin du document."""
    from generators.document_generator import (
        _docx_add_code,
        _docx_add_table,
        _docx_write_inlines,
    )
    from generators.markdown_document import (
        CodeBlock,
        Heading,
        ListBlock,
        Paragraph,
        Quote,
        Rule,
        Table,
    )

    for block in parse_markdown(content):
        if isinstance(block, Heading):
            _docx_write_inlines(
                document.add_heading("", level=min(block.level, 4)), block.inlines
            )
        elif isinstance(block, Paragraph):
            _docx_write_inlines(document.add_paragraph(), block.inlines)
        elif isinstance(block, ListBlock):
            style = "List Number" if block.ordered else "List Bullet"
            for item in block.items:
                _docx_write_inlines(document.add_paragraph(style=style), item.inlines)
        elif isinstance(block, Table):
            _docx_add_table(document, block)
        elif isinstance(block, CodeBlock):
            _docx_add_code(document, block)
        elif isinstance(block, Quote):
            _docx_write_inlines(document.add_paragraph(), block.inlines)
        elif isinstance(block, Rule):
            document.add_paragraph("─" * 40)


def _docx_replace_section(document, heading: str, content: str) -> int:
    """
    Remplace le contenu d'une section, de son titre au titre de même niveau.

    Les paragraphes du corps sont supprimés et le nouveau contenu Markdown est
    inséré à leur place, juste après le titre.
    """
    paragraphs = document.paragraphs
    start = None
    level = 0

    for index, paragraph in enumerate(paragraphs):
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading") and heading.lower() in paragraph.text.lower():
            start = index
            digits = "".join(char for char in style if char.isdigit())
            level = int(digits) if digits else 1
            break

    if start is None:
        return 0

    end = len(paragraphs)
    for index in range(start + 1, len(paragraphs)):
        style = (paragraphs[index].style.name or "").lower()
        if style.startswith("heading"):
            digits = "".join(char for char in style if char.isdigit())
            if (int(digits) if digits else 1) <= level:
                end = index
                break

    anchor = paragraphs[start]._p  # noqa: SLF001 - insertion positionnelle
    for paragraph in paragraphs[start + 1:end]:
        element = paragraph._p  # noqa: SLF001
        element.getparent().remove(element)

    # Le nouveau contenu est rendu en fin de document puis déplacé sous le titre.
    before = len(document.paragraphs)
    _docx_append_markdown(document, content)
    for paragraph in document.paragraphs[before:]:
        anchor.addnext(paragraph._p)  # noqa: SLF001
        anchor = paragraph._p  # noqa: SLF001

    return 1


# ══════════════════════════════════════════════════════════════════════════
#  Aides XLSX
# ══════════════════════════════════════════════════════════════════════════


def _resolve_sheet(workbook, name: str):
    """Retourne la feuille demandée, ou la feuille active si le nom est vide."""
    if not name:
        return workbook.active
    for sheet in workbook.worksheets:
        if sheet.title.lower() == name.lower():
            return sheet
    raise DocumentEditError(
        f"Feuille « {name} » introuvable. Feuilles disponibles : "
        f"{', '.join(workbook.sheetnames)}."
    )


# ══════════════════════════════════════════════════════════════════════════
#  Aides PPTX
# ══════════════════════════════════════════════════════════════════════════


def _pptx_iter_text_frames(presentation):
    """Itère sur tous les cadres de texte, cellules de tableau comprises."""
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                yield shape.text_frame
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame


def _pptx_replace_text(presentation, find: str, replace: str) -> int:
    """Remplace un texte dans toute la présentation, tableaux compris."""
    return sum(
        _replace_in_paragraph(paragraph, find, replace)
        for frame in _pptx_iter_text_frames(presentation)
        for paragraph in frame.paragraphs
    )


def _pptx_delete_paragraphs(presentation, needle: str) -> int:
    """Vide les paragraphes contenant un texte donné."""
    count = 0
    for frame in _pptx_iter_text_frames(presentation):
        for paragraph in frame.paragraphs:
            if needle in paragraph.text:
                for run in paragraph.runs:
                    run.text = ""
                count += 1
    return count


def _pptx_append_slide(presentation, title: str, content: str) -> None:
    """Ajoute une diapositive titre + puces à la fin de la présentation."""
    from pptx.util import Pt

    from generators.document_generator import _pptx_body_placeholder
    from generators.markdown_document import Heading, ListBlock, Paragraph

    layout = presentation.slide_layouts[1]
    slide = presentation.slides.add_slide(layout)
    if slide.shapes.title is not None:
        slide.shapes.title.text = title or "Nouvelle diapositive"

    body = _pptx_body_placeholder(slide)
    if body is None:
        return

    bullets: List[Tuple[str, int]] = []
    for block in parse_markdown(content):
        if isinstance(block, ListBlock):
            bullets.extend((item.text, item.level) for item in block.items)
        elif isinstance(block, (Paragraph, Heading)):
            bullets.append((block.text, 0))

    frame = body.text_frame
    frame.clear()
    for position, (text, level) in enumerate(bullets):
        paragraph = frame.paragraphs[0] if position == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.level = min(level, 4)
        paragraph.font.size = Pt(18 if level == 0 else 15)


def copy_to_outputs(path: str) -> Path:
    """
    Copie un document dans outputs/documents/ sans le modifier.

    Utile pour donner à l'utilisateur une copie de travail d'une pièce jointe.
    """
    source = Path(path).expanduser()
    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
    destination = DOCUMENTS_DIR / source.name
    counter = 2
    while destination.exists():
        destination = DOCUMENTS_DIR / f"{source.stem}_{counter}{source.suffix}"
        counter += 1
    shutil.copy2(source, destination)
    return destination
