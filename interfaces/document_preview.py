"""
Rendu HTML des documents pour le volet « Aperçu ».

Module 100% local, sans dépendance réseau, partagé par :
- le GUI desktop (volet de preview, cf. ``interfaces/gui/artifacts_panel.py``),
- le serveur Relay (HTML poussé chiffré vers la modale mobile).

Un document bureautique n'a pas de « code source » affichable : on en produit
une page HTML, mise en page comme une feuille de papier posée sur le fond
sombre du volet, cohérente avec le thème du GUI.

Le PDF fait exception : Edge et les navigateurs ont un lecteur PDF natif, donc
``needs_native_viewer()`` renvoie True et l'appelant sert le fichier tel quel.
"""

from __future__ import annotations

import html
from pathlib import Path
from typing import Any, Dict, List, Optional

from generators.markdown_document import (
    Block,
    CodeBlock,
    Heading,
    Inline,
    ListBlock,
    Paragraph,
    Quote,
    Rule,
    Table,
    parse_inlines,
    parse_markdown,
)

# Formats rendus en HTML par ce module.
_HTML_RENDERED = {".docx", ".xlsx", ".xlsm", ".xls", ".csv", ".pptx", ".potx",
                  ".md", ".markdown", ".txt", ".log", ".rst"}
# Formats affichés par le lecteur natif du moteur de rendu.
_NATIVE_VIEWER = {".pdf"}

# Libellé affiché dans l'en-tête de la feuille, par extension.
_FORMAT_LABELS = {
    ".docx": "Document Word", ".pdf": "Document PDF",
    ".pptx": "Présentation PowerPoint", ".potx": "Modèle PowerPoint",
    ".xlsx": "Classeur Excel", ".xlsm": "Classeur Excel", ".xls": "Classeur Excel",
    ".csv": "Données CSV", ".md": "Markdown", ".markdown": "Markdown",
    ".txt": "Texte", ".log": "Journal", ".rst": "reStructuredText",
}

# Styles de paragraphe Word correspondant à des listes.
_DOCX_LIST_STYLES = ("list bullet", "list number", "list paragraph")

# Au-delà, l'aperçu d'une feuille Excel est tronqué (le volet n'est pas un tableur).
_XLSX_MAX_ROWS = 200
_XLSX_MAX_COLUMNS = 40


# ── Page d'enveloppe ───────────────────────────────────────────────────────

_PAGE_TEMPLATE = """<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title}</title>
<style>
  :root {{
    --ink: #1a1a1a;
    --muted: #6b7280;
    --accent: #1f4e79;
    --rule: #e3e6ea;
    --sheet: #ffffff;
  }}
  html, body {{ margin: 0; padding: 0; }}
  body {{
    background: #212121;
    font-family: "Segoe UI", system-ui, -apple-system, sans-serif;
    padding: 20px 16px 48px;
    box-sizing: border-box;
  }}
  .sheet {{
    background: var(--sheet);
    color: var(--ink);
    max-width: 820px;
    margin: 0 auto;
    padding: 40px 48px 52px;
    border-radius: 6px;
    box-shadow: 0 6px 28px rgba(0, 0, 0, .45);
    line-height: 1.6;
    font-size: 15px;
  }}
  .doc-kind {{
    color: var(--muted);
    font-size: 11px;
    letter-spacing: .09em;
    text-transform: uppercase;
    margin-bottom: 6px;
  }}
  .doc-name {{
    color: var(--accent);
    font-size: 25px;
    font-weight: 650;
    margin: 0 0 22px;
    padding-bottom: 14px;
    border-bottom: 2px solid var(--rule);
    word-break: break-word;
  }}
  h1, h2, h3, h4, h5, h6 {{ color: var(--accent); line-height: 1.3; }}
  h1 {{ font-size: 22px; margin: 26px 0 10px; }}
  h2 {{ font-size: 19px; margin: 22px 0 9px; }}
  h3 {{ font-size: 16.5px; margin: 18px 0 8px; }}
  h4, h5, h6 {{ font-size: 15px; margin: 15px 0 7px; }}
  p {{ margin: 0 0 11px; }}
  ul, ol {{ margin: 0 0 12px; padding-left: 24px; }}
  li {{ margin-bottom: 4px; }}
  a {{ color: var(--accent); }}
  code {{
    background: #f2f3f5; border-radius: 3px; padding: 1px 5px;
    font-family: Consolas, "SF Mono", monospace; font-size: 13px;
  }}
  pre {{
    background: #f7f8fa; border: 1px solid var(--rule); border-left: 3px solid var(--accent);
    border-radius: 4px; padding: 12px 14px; overflow-x: auto;
    font-family: Consolas, "SF Mono", monospace; font-size: 12.5px; line-height: 1.5;
  }}
  pre code {{ background: none; padding: 0; font-size: inherit; }}
  blockquote {{
    margin: 12px 0; padding: 4px 0 4px 15px;
    border-left: 3px solid var(--accent); color: #4b5563; font-style: italic;
  }}
  hr {{ border: none; border-top: 1px solid var(--rule); margin: 22px 0; }}
  table {{ border-collapse: collapse; width: 100%; margin: 12px 0 18px; font-size: 13.5px; }}
  th, td {{ border: 1px solid #d8dde3; padding: 7px 10px; text-align: left; vertical-align: top; }}
  th {{ background: var(--accent); color: #fff; font-weight: 600; }}
  tbody tr:nth-child(even) {{ background: #f7f9fc; }}
  .sheet-title {{
    margin: 26px 0 8px; font-size: 13px; font-weight: 650; color: var(--accent);
    text-transform: uppercase; letter-spacing: .05em;
  }}
  .slide {{
    border: 1px solid var(--rule); border-radius: 6px;
    margin: 0 0 18px; overflow: hidden;
  }}
  .slide-head {{
    background: var(--accent); color: #fff;
    padding: 9px 14px; font-weight: 600; font-size: 15px;
  }}
  .slide-body {{ padding: 12px 16px; }}
  .slide-body ul {{ margin: 0; }}
  .slide-no {{ opacity: .65; font-weight: 400; font-size: 12px; margin-left: 6px; }}
  .slide-notes {{
    border-top: 1px dashed var(--rule); margin-top: 10px; padding-top: 8px;
    color: var(--muted); font-size: 12.5px;
  }}
  .empty {{ color: var(--muted); font-style: italic; }}
  .truncated {{
    color: var(--muted); font-size: 12.5px; font-style: italic; margin-top: 6px;
  }}
  @media (max-width: 640px) {{
    body {{ padding: 12px 8px 32px; }}
    .sheet {{ padding: 24px 20px 32px; font-size: 14.5px; }}
  }}
</style>
</head>
<body>
<div class="sheet">
{body}
</div>
</body>
</html>
"""


def wrap_preview_page(title: str, body: str) -> str:
    """
    Enveloppe un corps HTML dans la page « feuille » du volet d'aperçu.

    Args:
        title: Titre de la page (balise <title>)
        body: Fragment HTML déjà échappé

    Returns:
        Un document HTML complet et autonome.
    """
    return _PAGE_TEMPLATE.format(title=html.escape(title or "Aperçu"), body=body)


# ── Point d'entrée ─────────────────────────────────────────────────────────


def is_previewable(path: str) -> bool:
    """True si le document sait être affiché dans le volet d'aperçu."""
    suffix = Path(path).suffix.lower()
    return suffix in _HTML_RENDERED or suffix in _NATIVE_VIEWER


def needs_native_viewer(path: str) -> bool:
    """
    True si le fichier doit être servi tel quel au moteur de rendu.

    C'est le cas du PDF : Edge et les navigateurs l'affichent nativement, mieux
    que toute conversion HTML qu'on pourrait produire.
    """
    return Path(path).suffix.lower() in _NATIVE_VIEWER


def build_document_preview(path: str) -> str:
    """
    Construit la page d'aperçu HTML d'un document.

    Args:
        path: Chemin du document (docx, xlsx, pptx, csv, md, txt…)

    Returns:
        Un document HTML complet. En cas d'échec de lecture, une page
        expliquant le problème plutôt qu'une exception : le volet doit
        toujours afficher quelque chose.
    """
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    label = _FORMAT_LABELS.get(suffix, suffix.lstrip(".").upper() or "Fichier")
    header = (
        f'<div class="doc-kind">{html.escape(label)}</div>'
        f'<h1 class="doc-name">{html.escape(file_path.name)}</h1>'
    )

    try:
        if not file_path.exists():
            body = header + '<p class="empty">Fichier introuvable.</p>'
        elif suffix == ".docx":
            body = header + _render_docx(str(file_path))
        elif suffix in (".xlsx", ".xlsm", ".xls", ".csv"):
            body = header + _render_spreadsheet(str(file_path))
        elif suffix in (".pptx", ".potx"):
            body = header + _render_pptx(str(file_path))
        elif suffix == ".pdf":
            body = header + _render_pdf(str(file_path))
        elif suffix in (".md", ".markdown"):
            body = header + render_blocks_html(parse_markdown(_read_text(file_path)), "")
        else:
            body = header + f"<pre>{html.escape(_read_text(file_path))}</pre>"
    except Exception as exc:  # noqa: BLE001 - l'aperçu ne doit jamais faire tomber le volet
        body = header + (
            f'<p class="empty">Aperçu indisponible : {html.escape(str(exc))}</p>'
        )

    return wrap_preview_page(file_path.stem, body)


# ── Rendu depuis des blocs Markdown ────────────────────────────────────────


def render_blocks_html(blocks: List[Block], title: str = "") -> str:
    """
    Rend une liste de blocs (cf. ``generators/markdown_document``) en HTML.

    Args:
        blocks: Blocs issus de `parse_markdown`
        title: Titre optionnel ajouté en tête

    Returns:
        Un fragment HTML échappé, sans balise <html>.
    """
    parts: List[str] = []
    if title:
        parts.append(f'<h1 class="doc-name">{html.escape(title)}</h1>')

    index = 0
    while index < len(blocks):
        block = blocks[index]

        if isinstance(block, ListBlock):
            parts.append(_render_list(block))
        elif isinstance(block, Heading):
            level = min(block.level, 6)
            parts.append(f"<h{level}>{_render_inlines(block.inlines)}</h{level}>")
        elif isinstance(block, Paragraph):
            parts.append(f"<p>{_render_inlines(block.inlines)}</p>")
        elif isinstance(block, Table):
            parts.append(_render_markdown_table(block))
        elif isinstance(block, CodeBlock):
            parts.append(f"<pre><code>{html.escape(block.code)}</code></pre>")
        elif isinstance(block, Quote):
            parts.append(f"<blockquote>{_render_inlines(block.inlines)}</blockquote>")
        elif isinstance(block, Rule):
            parts.append("<hr>")
        index += 1

    return "\n".join(parts) or '<p class="empty">Document vide.</p>'


def _render_inlines(inlines: List[Inline]) -> str:
    """Convertit des fragments stylés en HTML échappé."""
    parts = []
    for text, styles, url in inlines:
        rendered = html.escape(text)
        if "code" in styles:
            rendered = f"<code>{rendered}</code>"
        if "bold" in styles:
            rendered = f"<strong>{rendered}</strong>"
        if "italic" in styles:
            rendered = f"<em>{rendered}</em>"
        if url:
            rendered = (
                f'<a href="{html.escape(url, quote=True)}" '
                f'target="_blank" rel="noopener noreferrer">{rendered}</a>'
            )
        parts.append(rendered)
    return "".join(parts)


def _render_list(block: ListBlock) -> str:
    """Rend une liste, en restituant les niveaux d'imbrication."""
    tag = "ol" if block.ordered else "ul"
    parts = [f"<{tag}>"]
    depth = 0
    for item in block.items:
        while depth < item.level:
            parts.append(f"<{tag}>")
            depth += 1
        while depth > item.level:
            parts.append(f"</{tag}>")
            depth -= 1
        parts.append(f"<li>{_render_inlines(item.inlines)}</li>")
    parts.extend(f"</{tag}>" for _ in range(depth + 1))
    return "".join(parts)


def _render_markdown_table(block: Table) -> str:
    """
    Rend un tableau issu du Markdown, en interprétant le gras/italique des cellules.

    Réservé aux tableaux Markdown : ceux lus dans un vrai fichier passent par
    _render_table, où un « 5 * 3 * 2 » ne doit pas devenir de l'italique.
    """
    parts = ["<table>"]
    if block.header:
        cells = "".join(
            f"<th>{_render_inlines(parse_inlines(value))}</th>" for value in block.header
        )
        parts.append(f"<thead><tr>{cells}</tr></thead>")
    parts.append("<tbody>")
    for row in block.rows:
        cells = "".join(f"<td>{_render_inlines(parse_inlines(value))}</td>" for value in row)
        parts.append(f"<tr>{cells}</tr>")
    parts.append("</tbody></table>")
    return "".join(parts)


def _render_table(header: List[str], rows: List[List[str]], note: str = "") -> str:
    """Rend un tableau HTML à partir d'un en-tête et de lignes de texte."""
    parts = ["<table>"]
    if header:
        cells = "".join(f"<th>{html.escape(str(value))}</th>" for value in header)
        parts.append(f"<thead><tr>{cells}</tr></thead>")
    parts.append("<tbody>")
    for row in rows:
        cells = "".join(f"<td>{html.escape(str(value))}</td>" for value in row)
        parts.append(f"<tr>{cells}</tr>")
    parts.append("</tbody></table>")
    if note:
        parts.append(f'<p class="truncated">{html.escape(note)}</p>')
    return "".join(parts)


# ── Rendu par format ───────────────────────────────────────────────────────


def _render_docx(path: str) -> str:
    """Rend un DOCX en HTML : titres, listes, paragraphes et tableaux."""
    from processors.docx_processor import DOCXProcessor

    result = DOCXProcessor().read_docx(path)
    if not result.get("success"):
        return f'<p class="empty">{html.escape(result.get("error", "Lecture impossible."))}</p>'

    content = result.get("content", {})
    parts: List[str] = []
    list_buffer: List[str] = []
    list_ordered = False

    def flush_list() -> None:
        if list_buffer:
            tag = "ol" if list_ordered else "ul"
            items = "".join(f"<li>{html.escape(text)}</li>" for text in list_buffer)
            parts.append(f"<{tag}>{items}</{tag}>")
            list_buffer.clear()

    for paragraph in content.get("paragraphs", []):
        text = (paragraph.get("text") or "").strip()
        style = (paragraph.get("style") or "Normal").lower()

        if not text:
            flush_list()
            continue

        if style.startswith("heading") or style in ("title", "subtitle"):
            flush_list()
            level = _docx_heading_level(style)
            parts.append(f"<h{level}>{html.escape(text)}</h{level}>")
        elif style.startswith(_DOCX_LIST_STYLES):
            ordered = style.startswith("list number")
            if list_buffer and ordered != list_ordered:
                flush_list()
            list_ordered = ordered
            list_buffer.append(text)
        elif "quote" in style:
            flush_list()
            parts.append(f"<blockquote>{html.escape(text)}</blockquote>")
        else:
            flush_list()
            parts.append(f"<p>{html.escape(text)}</p>")

    flush_list()

    for table in content.get("tables", []):
        rows = table.get("rows", [])
        if rows:
            parts.append(_render_table(rows[0], rows[1:]))

    return "\n".join(parts) or '<p class="empty">Document vide.</p>'


def _docx_heading_level(style: str) -> int:
    """Déduit le niveau de titre HTML d'un nom de style Word."""
    if style == "title":
        return 1
    if style == "subtitle":
        return 2
    digits = "".join(char for char in style if char.isdigit())
    return min(int(digits), 6) if digits else 2


def _render_spreadsheet(path: str) -> str:
    """Rend un classeur Excel ou un CSV : une section par feuille."""
    suffix = Path(path).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        sheets = _read_xlsx_sheets(path)
    else:
        sheets = _read_tabular_fallback(path)

    if not sheets:
        return '<p class="empty">Aucune donnée lisible.</p>'

    parts: List[str] = []
    for name, rows, truncated in sheets:
        parts.append(f'<div class="sheet-title">{html.escape(name)}</div>')
        if not rows:
            parts.append('<p class="empty">Feuille vide.</p>')
            continue
        note = "Aperçu tronqué — ouvrez le fichier pour tout voir." if truncated else ""
        parts.append(_render_table(rows[0], rows[1:], note))
    return "\n".join(parts)


def _read_xlsx_sheets(path: str) -> List[tuple]:
    """Lit un xlsx en (nom, lignes, tronqué) avec les valeurs calculées."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheets: List[tuple] = []
    try:
        for worksheet in workbook.worksheets:
            rows: List[List[str]] = []
            truncated = False
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if index >= _XLSX_MAX_ROWS:
                    truncated = True
                    break
                if len(row) > _XLSX_MAX_COLUMNS:
                    truncated = True
                cells = ["" if value is None else str(value) for value in row[:_XLSX_MAX_COLUMNS]]
                if any(cell.strip() for cell in cells):
                    rows.append(cells)
            sheets.append((worksheet.title, rows, truncated))
    finally:
        workbook.close()
    return sheets


def _read_tabular_fallback(path: str) -> List[tuple]:
    """Lit un CSV (ou un .xls) via ExcelProcessor, puis découpe en lignes."""
    from processors.excel_processor import ExcelProcessor

    result = ExcelProcessor().read_excel(path)
    if not result.get("success"):
        return []

    rows: List[List[str]] = []
    for line in (result.get("content") or "").splitlines():
        if not line.strip() or set(line.strip()) <= {"-", "+", "|", " "}:
            continue
        cells = [cell.strip() for cell in line.split("|")]
        # Les lignes produites par _rows_to_text sont encadrées de « | ».
        if cells and not cells[0]:
            cells = cells[1:]
        if cells and not cells[-1]:
            cells = cells[:-1]
        if cells:
            rows.append(cells)

    truncated = len(rows) > _XLSX_MAX_ROWS
    return [(Path(path).stem, rows[:_XLSX_MAX_ROWS], truncated)]


def _render_pptx(path: str) -> str:
    """Rend une présentation : une carte par diapositive."""
    from processors.pptx_processor import PPTXProcessor

    result = PPTXProcessor().read_pptx(path)
    if not result.get("success"):
        return f'<p class="empty">{html.escape(result.get("error", "Lecture impossible."))}</p>'

    slides = result.get("content", {}).get("slides", [])
    if not slides:
        return '<p class="empty">Présentation vide.</p>'

    parts: List[str] = []
    for slide in slides:
        number = slide.get("slide_number", 0)
        title = slide.get("title") or "(sans titre)"
        parts.append('<div class="slide">')
        parts.append(
            f'<div class="slide-head">{html.escape(title)}'
            f'<span class="slide-no">#{number}</span></div>'
        )
        parts.append('<div class="slide-body">')

        bullets = slide.get("bullets", [])
        if bullets:
            items = "".join(
                f"<li>{html.escape(text.strip())}</li>" for text in bullets if text.strip()
            )
            parts.append(f"<ul>{items}</ul>")

        for table in slide.get("tables", []):
            rows = table.get("rows", [])
            if rows:
                parts.append(_render_table(rows[0], rows[1:]))

        if not bullets and not slide.get("tables"):
            parts.append('<p class="empty">Diapositive sans texte.</p>')

        if slide.get("notes"):
            parts.append(
                f'<div class="slide-notes">Notes : {html.escape(slide["notes"])}</div>'
            )
        parts.append("</div></div>")

    return "\n".join(parts)


def _render_pdf(path: str) -> str:
    """
    Rend le texte extrait d'un PDF.

    Le desktop n'emprunte pas ce chemin (``needs_native_viewer`` renvoie True
    et Edge affiche le PDF lui-même) ; il sert à la modale mobile, dont
    l'iframe ``srcdoc`` ne sait pas afficher un PDF.
    """
    from processors.pdf_processor import PDFProcessor

    text = PDFProcessor().extract_text(path)
    if not text or not text.strip():
        return (
            '<p class="empty">Aucun texte extractible : PDF scanné, protégé '
            "ou constitué d'images.</p>"
        )

    parts = []
    for paragraph in text.split("\n\n"):
        cleaned = paragraph.strip()
        if cleaned:
            parts.append(f"<p>{html.escape(cleaned)}</p>")
    parts.append(
        '<p class="truncated">Texte extrait du PDF — la mise en page d\'origine '
        "n'est pas reproduite.</p>"
    )
    return "\n".join(parts)


def _read_text(path: Path) -> str:
    """Lit un fichier texte en UTF-8, avec repli latin-1."""
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="latin-1")


def document_summary(path: str) -> Dict[str, Any]:
    """
    Résume un document pour l'affichage (titre, format, taille).

    Utilisé par le volet d'aperçu et par le Relay pour étiqueter la modale.
    """
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    size: Optional[int] = None
    try:
        size = file_path.stat().st_size
    except OSError:
        pass
    return {
        "name": file_path.name,
        "title": file_path.stem,
        "format": suffix.lstrip("."),
        "label": _FORMAT_LABELS.get(suffix, suffix.lstrip(".").upper() or "Fichier"),
        "size": size,
        "native_viewer": needs_native_viewer(path),
    }
