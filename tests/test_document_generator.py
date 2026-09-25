"""
Tests du générateur de documents (generators/document_generator.py).

Génération RÉELLE (pas de mock) : chaque document est écrit dans un dossier
temporaire puis relu avec la bibliothèque du format, pour vérifier que la
structure Markdown est bien restituée (titres, listes, tableaux, gras) et pas
seulement que le fichier existe.

Le LLM n'est jamais sollicité : tous les tests fournissent `content`.
"""

import asyncio

import pytest

from generators.document_generator import DocumentGenerator
from generators.markdown_document import parse_markdown

MARKDOWN = """## Espèces principales

Les baleines sont des **mammifères marins** de l'ordre des *cétacés*.

- Baleine bleue
  - Sous-espèce antarctique
- Rorqual commun

1. Premier point
2. Second point

| Espèce | Taille | Poids |
|---|---|---|
| Bleue | 30 | 150 |
| À bosse | 16 | 30 |

> Elles communiquent par des chants.

## Conservation

Texte de la seconde section.

```python
print("ok")
```
"""


@pytest.fixture(name="generator")
def _generator():
    """Générateur sans LLM : les tests fournissent eux-mêmes le contenu."""
    return DocumentGenerator(llm=False)


def _build(generator, fmt, tmp_path, title="Les baleines", content=MARKDOWN):
    """Rend le Markdown de test dans `fmt` et retourne le chemin écrit."""
    path = tmp_path / f"doc.{fmt}"
    generator.build(parse_markdown(content), fmt, title, path)
    return path


# ── Détection du format ────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "query,expected",
    [
        ("génère moi un docx sur les baleines", "docx"),
        ("fais un document word", "docx"),
        ("un PDF de synthèse", "pdf"),
        ("crée une présentation sur le climat", "pptx"),
        ("un diaporama powerpoint", "pptx"),
        ("un tableur excel des ventes", "xlsx"),
        ("exporte en csv", "csv"),
        ("une note en markdown", "md"),
        ("un rapport sur les baleines", "docx"),  # défaut
        ("", "docx"),
    ],
)
def test_detect_format(query, expected):
    assert DocumentGenerator.detect_format(query) == expected


# ── DOCX ───────────────────────────────────────────────────────────────────


def test_docx_structure(generator, tmp_path):
    import docx

    path = _build(generator, "docx", tmp_path)
    document = docx.Document(str(path))
    styles = [p.style.name for p in document.paragraphs if p.text.strip()]
    texts = [p.text for p in document.paragraphs]

    assert "Title" in styles, "le titre du document doit utiliser le style Title"
    assert "Heading 2" in styles, "les ## doivent devenir des titres de niveau 2"
    assert "List Bullet" in styles
    assert "List Bullet 2" in styles, "l'imbrication de liste doit être conservée"
    assert "List Number" in styles
    assert "Les baleines" in texts

    # Tableau : en-tête + deux lignes de données
    assert len(document.tables) == 1
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["Espèce", "Taille", "Poids"]
    assert len(table.rows) == 3

    # Le gras du Markdown doit exister comme run gras
    assert any(run.bold for p in document.paragraphs for run in p.runs)


def test_docx_adds_toc_from_three_headings(generator, tmp_path):
    import docx

    content = "## Un\n\nA\n\n## Deux\n\nB\n\n## Trois\n\nC"
    path = _build(generator, "docx", tmp_path, content=content)
    document = docx.Document(str(path))

    assert "Sommaire" in [p.text for p in document.paragraphs]
    assert "TOC" in document.element.xml, "le champ TOC doit être inséré"


def test_docx_footer_has_page_number(generator, tmp_path):
    import docx

    path = _build(generator, "docx", tmp_path)
    footer = docx.Document(str(path)).sections[0].footer
    assert "PAGE" in footer.paragraphs[0]._p.xml
    assert "Page" in footer.paragraphs[0].text


def test_docx_without_enough_headings_has_no_toc(generator, tmp_path):
    import docx

    path = _build(generator, "docx", tmp_path, content="## Une seule section\n\nTexte.")
    document = docx.Document(str(path))
    assert "Sommaire" not in [p.text for p in document.paragraphs]


# ── PDF ────────────────────────────────────────────────────────────────────


def test_pdf_is_readable_and_keeps_text(generator, tmp_path):
    from processors.pdf_processor import PDFProcessor

    path = _build(generator, "pdf", tmp_path)
    assert path.stat().st_size > 1000

    text = PDFProcessor().extract_text(str(path))
    assert "Les baleines" in text
    assert "Espèces principales" in text
    assert "Conservation" in text
    assert "Rorqual commun" in text


# ── PPTX ───────────────────────────────────────────────────────────────────


def test_pptx_one_slide_per_section(generator, tmp_path):
    from pptx import Presentation

    path = _build(generator, "pptx", tmp_path)
    presentation = Presentation(str(path))
    slides = list(presentation.slides)
    titles = [s.shapes.title.text for s in slides if s.shapes.title is not None]

    assert titles[0] == "Les baleines", "la première diapo est la couverture"
    assert any(t.startswith("Espèces principales") for t in titles)
    assert "Conservation" in titles
    # Un tableau Markdown devient un vrai tableau PowerPoint
    assert any(shape.has_table for slide in slides for shape in slide.shapes)


def test_pptx_splits_long_sections(generator, tmp_path):
    from pptx import Presentation

    bullets = "\n".join(f"- Point {i}" for i in range(25))
    path = _build(generator, "pptx", tmp_path, content=f"## Longue section\n\n{bullets}")
    titles = [
        s.shapes.title.text
        for s in Presentation(str(path)).slides
        if s.shapes.title is not None
    ]
    continuations = [t for t in titles if t.startswith("Longue section")]
    assert len(continuations) >= 3, "25 puces doivent déborder sur plusieurs diapos"
    assert "Longue section (2)" in titles


def test_pptx_heading_without_bullets_creates_no_empty_slide(generator, tmp_path):
    from pptx import Presentation

    content = "## Titre seul\n\n| A | B |\n|---|---|\n| 1 | 2 |"
    path = _build(generator, "pptx", tmp_path, content=content)
    slides = list(Presentation(str(path)).slides)
    # Couverture + diapo du tableau, pas de diapo vide intercalée
    assert len(slides) == 2
    assert any(shape.has_table for shape in slides[1].shapes)


# ── XLSX ───────────────────────────────────────────────────────────────────


def test_xlsx_table_becomes_a_sheet(generator, tmp_path):
    from openpyxl import load_workbook

    path = _build(generator, "xlsx", tmp_path)
    workbook = load_workbook(str(path))

    assert "Notes" in workbook.sheetnames
    sheet = workbook["Espèces principales"]
    assert [cell.value for cell in sheet[1]] == ["Espèce", "Taille", "Poids"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None


def test_xlsx_numeric_cells_are_numbers(generator, tmp_path):
    from openpyxl import load_workbook

    path = _build(generator, "xlsx", tmp_path)
    sheet = load_workbook(str(path))["Espèces principales"]
    assert sheet["B2"].value == 30
    assert isinstance(sheet["C2"].value, int)
    assert sheet["A2"].value == "Bleue"


def test_xlsx_sheet_names_are_deduplicated_and_sanitized(generator, tmp_path):
    from openpyxl import load_workbook

    content = (
        "## Ventes/2024\n\n| A |\n|---|\n| 1 |\n\n"
        "## Ventes/2024\n\n| B |\n|---|\n| 2 |\n"
    )
    path = _build(generator, "xlsx", tmp_path, content=content)
    names = [n for n in load_workbook(str(path)).sheetnames if n != "Notes"]
    assert len(names) == 2 and len(set(names)) == 2
    assert all("/" not in name for name in names)


# ── Markdown dans les cellules de tableau ──────────────────────────────────
#
# Régression : les LLM mettent souvent en gras la première colonne
# (« | **Taille** | 30 m | ») ; les astérisques apparaissaient tels quels dans
# le document, car les cellules étaient écrites sans interprétation.

_TABLE_MD = (
    "## Récapitulatif\n\n"
    "| Critère | Valeur |\n|---|---|\n"
    "| **Taille** | Grande |\n"
    "| **Poids** | *environ* 150 |\n"
    "| Rang | 30 |\n"
)


def test_docx_table_cells_render_markdown(generator, tmp_path):
    import docx

    path = _build(generator, "docx", tmp_path, content=_TABLE_MD)
    table = docx.Document(str(path)).tables[0]
    assert table.rows[1].cells[0].text == "Taille"
    runs = [run for run in table.rows[1].cells[0].paragraphs[0].runs if run.text]
    assert runs and all(run.bold for run in runs)
    italic = [run.text for run in table.rows[2].cells[1].paragraphs[0].runs if run.italic]
    assert italic == ["environ"]
    assert "*" not in "".join(cell.text for row in table.rows for cell in row.cells)


def test_pdf_table_cells_have_no_markers(generator, tmp_path):
    from processors.pdf_processor import PDFProcessor

    text = PDFProcessor().extract_text(str(_build(generator, "pdf", tmp_path, content=_TABLE_MD)))
    assert "Taille" in text and "*" not in text


def test_pptx_table_cells_render_markdown(generator, tmp_path):
    from pptx import Presentation

    path = _build(generator, "pptx", tmp_path, content=_TABLE_MD)
    cells = [
        cell
        for slide in Presentation(str(path)).slides
        for shape in slide.shapes if shape.has_table
        for row in shape.table.rows for cell in row.cells
    ]
    taille = next(cell for cell in cells if cell.text == "Taille")
    assert all(run.font.bold for run in taille.text_frame.paragraphs[0].runs)
    assert not any("*" in cell.text for cell in cells)


def test_xlsx_table_cells_are_clean_and_bold(generator, tmp_path):
    from openpyxl import load_workbook

    path = _build(generator, "xlsx", tmp_path, content=_TABLE_MD)
    sheet = load_workbook(str(path))["Récapitulatif"]
    assert sheet["A2"].value == "Taille"
    assert sheet["A2"].font.b is True, "une cellule entièrement en gras doit l'être dans Excel"
    assert sheet["A4"].font.b is not True
    assert sheet["B4"].value == 30, "le nombre doit rester un nombre une fois nettoyé"


@pytest.mark.parametrize("fmt", ["csv", "txt"])
def test_plain_text_tables_have_no_markers(generator, tmp_path, fmt):
    encoding = "utf-8-sig" if fmt == "csv" else "utf-8"
    text = _build(generator, fmt, tmp_path, content=_TABLE_MD).read_text(encoding=encoding)
    assert "Taille" in text and "*" not in text


# ── Formats texte ──────────────────────────────────────────────────────────


def test_markdown_roundtrip_keeps_structure(generator, tmp_path):
    path = _build(generator, "md", tmp_path)
    text = path.read_text(encoding="utf-8")
    assert text.startswith("# Les baleines")
    assert "## Espèces principales" in text
    assert "| Espèce | Taille | Poids |" in text
    assert "```python" in text


def test_txt_renders_table_as_aligned_columns(generator, tmp_path):
    path = _build(generator, "txt", tmp_path)
    text = path.read_text(encoding="utf-8")
    assert "Les baleines" in text
    assert "Espèce" in text and "Taille" in text
    assert "- Baleine bleue" in text


def test_html_is_a_standalone_page(generator, tmp_path):
    path = _build(generator, "html", tmp_path)
    text = path.read_text(encoding="utf-8")
    assert text.lstrip().startswith("<!DOCTYPE html>")
    assert "<table>" in text and "<strong>" in text


def test_csv_exports_the_first_table(generator, tmp_path):
    path = _build(generator, "csv", tmp_path)
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    assert lines[0] == "Espèce;Taille;Poids"
    assert "Bleue;30;150" in lines


def test_unknown_format_is_refused(generator, tmp_path):
    with pytest.raises(ValueError):
        generator.build([], "odt", "T", tmp_path / "x.odt")


# ── generate_document (bout en bout, sans LLM) ─────────────────────────────


def test_generate_document_returns_metadata(generator):
    result = asyncio.run(
        generator.generate_document(
            "un rapport sur les baleines",
            fmt="docx",
            title="Les baleines",
            content=MARKDOWN,
            filename="pytest_baleines",
        )
    )
    assert result["success"] is True
    assert result["file_name"] == "pytest_baleines.docx"
    assert result["format"] == "docx"
    assert result["size"] > 0

    from pathlib import Path

    path = Path(result["file_path"])
    assert path.exists()
    path.unlink()


def test_generate_document_returns_the_outline(generator):
    """Le plan permet à la synthèse de présenter le document sans l'inventer."""
    result = asyncio.run(
        generator.generate_document(
            "un rapport",
            fmt="md",
            title="Les baleines",
            content="## Introduction\n\nA\n\n### Détail\n\nB\n\n## Conservation\n\nC",
            filename="pytest_outline",
        )
    )
    try:
        assert result["success"] is True
        # Niveaux 1 et 2 seulement : un plan, pas une table des matières complète
        assert result["outline"] == ["Introduction", "Conservation"]
    finally:
        from pathlib import Path

        Path(result["file_path"]).unlink(missing_ok=True)


def test_generate_document_without_llm_reports_the_cause(generator):
    """Sans Ollama et sans contenu fourni, l'échec doit être explicite."""
    result = asyncio.run(generator.generate_document("un sujet quelconque", fmt="docx"))
    assert result["success"] is False
    assert "Ollama" in result["error"]


def test_generate_document_derives_title_from_content(generator, tmp_path):
    result = asyncio.run(
        generator.generate_document(
            "peu importe",
            fmt="md",
            content="# Titre déduit\n\nCorps.",
            filename=str(tmp_path / "derive"),
        )
    )
    assert result["success"] is True
    assert result["title"] == "Titre déduit"
    # Le titre ne doit pas être dupliqué dans le corps
    from pathlib import Path

    text = Path(result["file_path"]).read_text(encoding="utf-8")
    assert text.count("Titre déduit") == 1
    Path(result["file_path"]).unlink()
