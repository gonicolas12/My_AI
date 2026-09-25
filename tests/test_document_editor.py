"""
Tests de l'éditeur de documents (generators/document_editor.py).

Édition RÉELLE sur des fichiers écrits par DocumentGenerator. Le point
critique vérifié partout : **le fichier source n'est jamais modifié**.
"""

import pytest

from generators.document_editor import DocumentEditor, _normalize_operations
from generators.document_generator import DOCUMENTS_DIR, DocumentGenerator
from generators.markdown_document import parse_markdown

MARKDOWN = """## Espèces

Le **rorqual commun** vit en haute mer.

- Baleine bleue
- Rorqual commun

| Espèce | Taille |
|---|---|
| Bleue | 30 |
| Bosse | 16 |

## Conservation

Statut préoccupant.
"""


@pytest.fixture(name="generator")
def _generator():
    return DocumentGenerator(llm=False)


@pytest.fixture(name="editor")
def _editor(generator):
    return DocumentEditor(generator=generator)


@pytest.fixture(name="source_factory")
def _source_factory(generator, tmp_path):
    """Écrit un document source dans tmp_path et retourne son chemin."""

    def _make(fmt, content=MARKDOWN, title="Les baleines"):
        path = tmp_path / f"source.{fmt}"
        generator.build(parse_markdown(content), fmt, title, path)
        return path

    return _make


@pytest.fixture(autouse=True)
def _cleanup_outputs():
    """Retire les copies produites dans outputs/documents/ après chaque test."""
    before = set(DOCUMENTS_DIR.glob("*")) if DOCUMENTS_DIR.exists() else set()
    yield
    if DOCUMENTS_DIR.exists():
        for path in set(DOCUMENTS_DIR.glob("*")) - before:
            try:
                path.unlink()
            except OSError:
                pass


def _snapshot(path):
    """Empreinte (taille, mtime, octets) pour détecter toute écriture."""
    return path.stat().st_size, path.stat().st_mtime_ns, path.read_bytes()


# ── Invariant : la source n'est jamais touchée ─────────────────────────────


@pytest.mark.parametrize("fmt", ["docx", "xlsx", "pptx", "md"])
def test_source_file_is_never_modified(editor, source_factory, fmt):
    source = source_factory(fmt)
    before = _snapshot(source)

    operations = {
        "docx": [{"action": "replace_text", "find": "Rorqual", "replace": "Baleinoptère"}],
        "xlsx": [{"action": "set_cell", "sheet": "Espèces", "cell": "B2", "value": 33}],
        "pptx": [{"action": "replace_text", "find": "Bleue", "replace": "Azur"}],
        "md": [{"action": "replace_text", "find": "Rorqual", "replace": "Baleinoptère"}],
    }[fmt]

    result = editor.edit(str(source), operations)
    assert result["success"] is True, result.get("error")
    assert _snapshot(source) == before, "le fichier source a été modifié"
    assert result["file_path"] != str(source)
    assert str(DOCUMENTS_DIR) in result["file_path"]


def test_copy_is_never_overwritten(editor, source_factory):
    source = source_factory("md")
    operations = [{"action": "replace_text", "find": "Espèces", "replace": "Taxons"}]

    first = editor.edit(str(source), operations)
    second = editor.edit(str(source), operations)

    assert first["success"] and second["success"]
    assert first["file_path"] != second["file_path"], "la 1re copie a été écrasée"


# ── DOCX ───────────────────────────────────────────────────────────────────


def test_docx_replace_text_preserves_formatting(editor, source_factory):
    import docx

    source = source_factory("docx")
    result = editor.edit(
        str(source),
        [{"action": "replace_text", "find": "rorqual commun", "replace": "petit rorqual"}],
    )
    assert result["success"] and result["applied"] >= 1

    document = docx.Document(result["file_path"])
    text = "\n".join(p.text for p in document.paragraphs)
    assert "petit rorqual" in text
    assert "rorqual commun" not in text
    # Le texte remplacé était en gras : le run doit le rester
    bold_runs = [r.text for p in document.paragraphs for r in p.runs if r.bold]
    assert any("petit rorqual" in t for t in bold_runs)


def test_docx_append_markdown_adds_structured_content(editor, source_factory):
    import docx

    source = source_factory("docx")
    result = editor.edit(
        str(source),
        [{"action": "append_markdown", "content": "## Annexe\n\n- Point ajouté"}],
    )
    document = docx.Document(result["file_path"])
    styles = {p.style.name: p.text for p in document.paragraphs if p.text.strip()}
    assert "Annexe" in styles.get("Heading 2", "") or "Annexe" in document.paragraphs[-2].text
    assert any("Point ajouté" in p.text for p in document.paragraphs)


def test_docx_delete_paragraph(editor, source_factory):
    import docx

    source = source_factory("docx")
    result = editor.edit(
        str(source), [{"action": "delete_paragraph", "contains": "Statut préoccupant"}]
    )
    assert result["applied"] == 1
    text = "\n".join(p.text for p in docx.Document(result["file_path"]).paragraphs)
    assert "Statut préoccupant" not in text


def test_docx_replace_section(editor, source_factory):
    import docx

    source = source_factory("docx")
    result = editor.edit(
        str(source),
        [{"action": "replace_section", "heading": "Conservation", "content": "Texte neuf."}],
    )
    assert result["success"] and result["applied"] == 1
    text = "\n".join(p.text for p in docx.Document(result["file_path"]).paragraphs)
    assert "Texte neuf." in text
    assert "Statut préoccupant" not in text
    assert "Conservation" in text, "le titre de section doit être conservé"


def test_docx_replace_section_missing_is_reported(editor, source_factory):
    source = source_factory("docx")
    result = editor.edit(
        str(source),
        [
            {"action": "replace_section", "heading": "Inexistante", "content": "x"},
            {"action": "replace_text", "find": "Espèces", "replace": "Taxons"},
        ],
    )
    assert result["success"] is True
    assert any("Inexistante" in note for note in result["notes"])


# ── XLSX ───────────────────────────────────────────────────────────────────


def test_xlsx_set_cell_and_append_row(editor, source_factory):
    from openpyxl import load_workbook

    source = source_factory("xlsx")
    result = editor.edit(
        str(source),
        [
            {"action": "set_cell", "sheet": "Espèces", "cell": "B2", "value": 33},
            {"action": "append_row", "sheet": "Espèces", "values": ["Grise", 14]},
        ],
    )
    assert result["success"] and result["applied"] == 2

    sheet = load_workbook(result["file_path"])["Espèces"]
    assert sheet["B2"].value == 33
    assert sheet.cell(row=sheet.max_row, column=1).value == "Grise"


def test_xlsx_unknown_sheet_is_refused_with_a_useful_message(editor, source_factory):
    source = source_factory("xlsx")
    result = editor.edit(
        str(source), [{"action": "set_cell", "sheet": "Néant", "cell": "A1", "value": 1}]
    )
    assert result["success"] is False
    assert "Néant" in result["error"] and "disponibles" in result["error"]


# ── PPTX ───────────────────────────────────────────────────────────────────


def test_pptx_replace_text_reaches_table_cells(editor, source_factory):
    from pptx import Presentation

    source = source_factory("pptx")
    result = editor.edit(
        str(source), [{"action": "replace_text", "find": "Bleue", "replace": "Azur"}]
    )
    assert result["success"] and result["applied"] >= 1

    cells = [
        cell.text
        for slide in Presentation(result["file_path"]).slides
        for shape in slide.shapes
        if shape.has_table
        for row in shape.table.rows
        for cell in row.cells
    ]
    assert "Azur" in cells and "Bleue" not in cells


def test_pptx_append_slide(editor, source_factory):
    from pptx import Presentation

    source = source_factory("pptx")
    before = len(list(Presentation(str(source)).slides))
    result = editor.edit(
        str(source),
        [{"action": "append_slide", "title": "Merci", "content": "- Questions\n- Contact"}],
    )
    slides = list(Presentation(result["file_path"]).slides)
    assert len(slides) == before + 1
    assert slides[-1].shapes.title.text == "Merci"


# ── PDF ────────────────────────────────────────────────────────────────────


def test_pdf_is_regenerated_and_says_so(editor, source_factory):
    from processors.pdf_processor import PDFProcessor

    source = source_factory("pdf")
    before = _snapshot(source)
    result = editor.edit(
        str(source), [{"action": "replace_text", "find": "Conservation", "replace": "Protection"}]
    )

    assert result["success"] is True
    assert _snapshot(source) == before
    assert any("pas modifiable en place" in note for note in result["notes"])

    text = PDFProcessor().extract_text(result["file_path"])
    assert "Protection" in text and "Conservation" not in text


# ── Erreurs et normalisation ───────────────────────────────────────────────


def test_missing_file_is_reported(editor):
    result = editor.edit("nulle_part/absent.docx", [{"action": "replace_text", "find": "a"}])
    assert result["success"] is False
    assert "introuvable" in result["error"].lower()


def test_unsupported_format_is_reported(editor, tmp_path):
    path = tmp_path / "image.png"
    path.write_bytes(b"\x89PNG\r\n")
    result = editor.edit(str(path), [{"action": "replace_text", "find": "a"}])
    assert result["success"] is False
    assert "non modifiable" in result["error"]


def test_no_operation_is_reported(editor, source_factory):
    source = source_factory("md")
    assert editor.edit(str(source), [])["success"] is False
    assert editor.edit(str(source), [{"action": "inconnue"}])["success"] is False


def test_unmatched_replacement_is_reported(editor, source_factory):
    source = source_factory("md")
    result = editor.edit(
        str(source), [{"action": "replace_text", "find": "texte absent", "replace": "x"}]
    )
    assert result["success"] is False
    assert "introuvable" in result["error"]


def test_missing_required_field_is_reported(editor, source_factory):
    source = source_factory("md")
    result = editor.edit(str(source), [{"action": "replace_text", "replace": "x"}])
    assert result["success"] is False
    assert "find" in result["error"]


def test_operation_aliases_are_normalized():
    """Les modèles locaux nomment les champs de façons variées."""
    operations = _normalize_operations(
        {"action": "remplacer", "old_text": "a", "new_text": "b"}
    )
    assert operations == [{"action": "replace_text", "find": "a", "replace": "b"}]

    assert _normalize_operations([{"action": "add_slide", "titre": "T", "text": "C"}]) == [
        {"action": "append_slide", "title": "T", "content": "C"}
    ]
    # Entrées inexploitables écartées
    assert _normalize_operations(["pas un dict", {"action": "n_existe_pas"}]) == []
    assert _normalize_operations(None) == []


def test_output_name_is_respected(editor, source_factory):
    source = source_factory("md")
    result = editor.edit(
        str(source),
        [{"action": "replace_text", "find": "Espèces", "replace": "Taxons"}],
        output_name="mon_resultat.md",
    )
    assert result["file_name"] == "mon_resultat.md"


def test_copy_to_outputs_leaves_the_source_alone(source_factory):
    from generators.document_editor import copy_to_outputs

    source = source_factory("md")
    before = _snapshot(source)
    copy = copy_to_outputs(str(source))
    try:
        assert copy.exists()
        assert copy.read_bytes() == source.read_bytes()
        assert _snapshot(source) == before
    finally:
        copy.unlink(missing_ok=True)
