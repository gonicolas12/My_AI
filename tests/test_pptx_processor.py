"""
Tests du processeur PPTX (processors/pptx_processor.py).

Lecture RÉELLE d'une présentation construite à la volée avec python-pptx,
puis d'une présentation produite par DocumentGenerator — c'est le chemin
qu'emprunte une pièce jointe .pptx du chat.
"""

import pytest
from pptx import Presentation
from pptx.util import Inches

from generators.document_generator import DocumentGenerator
from generators.markdown_document import parse_markdown
from processors.pptx_processor import PPTXProcessor


@pytest.fixture(name="processor")
def _processor():
    return PPTXProcessor()


@pytest.fixture(name="deck")
def _deck(tmp_path):
    """Présentation de test : titre, puces à deux niveaux, tableau, notes."""
    presentation = Presentation()

    cover = presentation.slides.add_slide(presentation.slide_layouts[0])
    cover.shapes.title.text = "Les baleines"
    cover.placeholders[1].text = "Sous-titre"

    content = presentation.slides.add_slide(presentation.slide_layouts[1])
    content.shapes.title.text = "Espèces"
    frame = content.placeholders[1].text_frame
    frame.text = "Baleine bleue"
    nested = frame.add_paragraph()
    nested.text = "Sous-espèce antarctique"
    nested.level = 1
    content.notes_slide.notes_text_frame.text = "Insister sur la taille."

    table_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table_slide.shapes.title.text = "Mesures"
    table = table_slide.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(6), Inches(1)
    ).table
    table.cell(0, 0).text = "Espèce"
    table.cell(0, 1).text = "Taille"
    table.cell(1, 0).text = "Bleue"
    table.cell(1, 1).text = "30 m"

    path = tmp_path / "deck.pptx"
    presentation.save(str(path))
    return path


# ── Lecture ────────────────────────────────────────────────────────────────


def test_read_pptx_returns_slides(processor, deck):
    result = processor.read_pptx(str(deck))
    assert result["success"] is True

    slides = result["content"]["slides"]
    assert len(slides) == 3
    assert [s["slide_number"] for s in slides] == [1, 2, 3]
    assert [s["title"] for s in slides] == ["Les baleines", "Espèces", "Mesures"]


def test_bullets_keep_their_indentation(processor, deck):
    slides = processor.read_pptx(str(deck))["content"]["slides"]
    bullets = slides[1]["bullets"]
    assert "Baleine bleue" in bullets
    assert any(b.startswith("    ") and "antarctique" in b for b in bullets), (
        "le niveau d'imbrication doit apparaître dans l'indentation"
    )


def test_title_is_not_repeated_in_bullets(processor, deck):
    slides = processor.read_pptx(str(deck))["content"]["slides"]
    assert "Espèces" not in slides[1]["bullets"]


def test_tables_and_notes_are_extracted(processor, deck):
    slides = processor.read_pptx(str(deck))["content"]["slides"]
    assert slides[1]["notes"] == "Insister sur la taille."
    assert slides[2]["tables"] == [{"rows": [["Espèce", "Taille"], ["Bleue", "30 m"]]}]


def test_properties_report_the_slide_count(processor, deck):
    properties = processor.read_pptx(str(deck))["content"]["properties"]
    assert properties["slide_count"] == 3


def test_text_representation_is_usable_as_context(processor, deck):
    text = processor.read_pptx(str(deck))["content"]["text"]
    assert "--- Diapositive 1 ---" in text
    assert "Les baleines" in text
    assert "Baleine bleue" in text
    assert "Espèce | Taille" in text
    assert "[Notes] Insister sur la taille." in text


# ── Interface commune aux processeurs ──────────────────────────────────────


def test_extract_text(processor, deck):
    result = processor.extract_text(str(deck))
    assert result["success"] is True
    assert result["metadata"]["slides"] == 3
    assert "Les baleines" in result["content"]


def test_extract_text_from_pptx(processor, deck):
    assert "Les baleines" in processor.extract_text_from_pptx(str(deck))


def test_template_potx_is_readable(processor, deck, tmp_path):
    """
    Un .potx ne diffère d'un .pptx que par le type de contenu de sa partie
    principale, que python-pptx refuse : le modèle doit tout de même se lire.
    """
    import zipfile

    main = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
    template = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
    potx = tmp_path / "modele.potx"
    with zipfile.ZipFile(deck) as source, zipfile.ZipFile(potx, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(main.encode(), template.encode())
            target.writestr(item, data)
    before = potx.read_bytes()

    result = processor.read_pptx(str(potx))
    assert result["success"] is True, result.get("error")
    assert [s["title"] for s in result["content"]["slides"]] == ["Les baleines", "Espèces", "Mesures"]
    # Le type de contenu est corrigé en mémoire, jamais dans le fichier
    assert potx.read_bytes() == before


def test_is_supported(processor):
    assert processor.is_supported("x.pptx") is True
    assert processor.is_supported("X.PPTX") is True
    assert processor.is_supported("x.potx") is True
    # .ppt (binaire pré-2007) n'est pas lisible par python-pptx
    assert processor.is_supported("x.ppt") is False
    assert processor.is_supported("x.docx") is False


# ── Erreurs ────────────────────────────────────────────────────────────────


def test_missing_file_is_reported(processor):
    result = processor.read_pptx("nulle_part/absent.pptx")
    assert result["success"] is False
    assert result["error_type"] == "FILE_NOT_FOUND"


def test_corrupt_file_is_reported(processor, tmp_path):
    path = tmp_path / "casse.pptx"
    path.write_bytes(b"ceci n'est pas une presentation")
    result = processor.read_pptx(str(path))
    assert result["success"] is False
    assert result["error_type"] == "PROCESSING_ERROR"


def test_extract_text_from_pptx_raises_on_error(processor):
    with pytest.raises(Exception, match="extraction PPTX"):
        processor.extract_text_from_pptx("nulle_part/absent.pptx")


# ── Aller-retour avec le générateur ────────────────────────────────────────


def test_generated_deck_is_readable(processor, tmp_path):
    """Ce que My_AI génère doit pouvoir être relu comme pièce jointe."""
    markdown = "## Section A\n\n- Un\n- Deux\n\n## Section B\n\n- Trois"
    path = tmp_path / "genere.pptx"
    DocumentGenerator(llm=False).build(parse_markdown(markdown), "pptx", "Titre", path)

    result = processor.read_pptx(str(path))
    assert result["success"] is True
    titles = [s["title"] for s in result["content"]["slides"]]
    assert "Titre" in titles
    assert "Section A" in titles and "Section B" in titles


# ── Intégration avec le pipeline de pièces jointes ────────────────────────


def test_file_processor_handles_pptx(deck):
    from utils.file_processor import FileProcessor

    file_processor = FileProcessor()
    assert file_processor.is_supported(str(deck)) is True

    result = file_processor.process_file(str(deck))
    assert result.get("error") is None
    assert result["type"] == "pptx"
    assert result["slides"] == 3
    assert "Les baleines" in result["content"]
