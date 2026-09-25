"""
Tests de l'aperçu HTML des documents (interfaces/document_preview.py) et des
artifacts « document » (interfaces/artifacts.py).

C'est ce qui alimente le volet latéral du GUI et la modale du Relay mobile :
le rendu doit être une page autonome, correctement échappée, et ne jamais
lever d'exception — un aperçu raté ne doit pas faire tomber le volet.
"""

import pytest

from generators.document_generator import DocumentGenerator
from generators.markdown_document import parse_markdown
from interfaces.artifacts import (
    artifact_from_document,
    artifacts_from_documents,
    build_preview_document,
    detect_artifacts,
    write_artifact_html,
)
from interfaces.document_preview import (
    build_document_preview,
    document_summary,
    is_previewable,
    needs_native_viewer,
    render_blocks_html,
)

MARKDOWN = """## Espèces

Les baleines sont des **mammifères marins**.

- Baleine bleue
  - Sous-espèce antarctique
- Rorqual commun

| Espèce | Taille |
|---|---|
| Bleue | 30 |

> Une citation.

```python
print("ok")
```
"""


@pytest.fixture(name="documents")
def _documents(tmp_path):
    """Un document de test par format, écrit dans tmp_path."""
    generator = DocumentGenerator(llm=False)
    blocks = parse_markdown(MARKDOWN)
    paths = {}
    for fmt in ("docx", "pdf", "pptx", "xlsx", "csv", "md", "txt"):
        path = tmp_path / f"doc.{fmt}"
        generator.build(blocks, fmt, "Les baleines", path)
        paths[fmt] = path
    return paths


# ── Classification ─────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "name,previewable", [
        ("a.docx", True), ("a.pdf", True), ("a.pptx", True), ("a.xlsx", True),
        ("a.csv", True), ("a.md", True), ("a.txt", True),
        ("a.exe", False), ("a.zip", False), ("a.png", False),
    ],
)
def test_is_previewable(name, previewable):
    assert is_previewable(name) is previewable


def test_only_pdf_uses_the_native_viewer():
    assert needs_native_viewer("a.pdf") is True
    assert needs_native_viewer("a.docx") is False
    assert needs_native_viewer("a.pptx") is False


# ── Rendu par format ───────────────────────────────────────────────────────


@pytest.mark.parametrize("fmt", ["docx", "pdf", "pptx", "xlsx", "csv", "md", "txt"])
def test_preview_is_a_standalone_page(documents, fmt):
    html = build_document_preview(str(documents[fmt]))
    assert html.lstrip().startswith("<!DOCTYPE html>")
    assert '<div class="sheet">' in html
    assert html.rstrip().endswith("</html>")
    assert f"doc.{fmt}" in html, "le nom du fichier doit apparaître en en-tête"


def test_docx_preview_keeps_structure(documents):
    html = build_document_preview(str(documents["docx"]))
    assert "Espèces" in html
    assert "<ul>" in html and "Baleine bleue" in html
    assert "<table>" in html and "<th>Espèce</th>" in html


def test_pptx_preview_shows_one_card_per_slide(documents):
    html = build_document_preview(str(documents["pptx"]))
    assert html.count('<div class="slide">') >= 2
    assert "slide-head" in html
    assert "Espèces" in html


def test_xlsx_preview_shows_each_sheet(documents):
    html = build_document_preview(str(documents["xlsx"]))
    assert "sheet-title" in html
    assert "<table>" in html
    assert "Espèces" in html


def test_pdf_preview_renders_extracted_text(documents):
    html = build_document_preview(str(documents["pdf"]))
    assert "Les baleines" in html
    assert "mise en page d'origine" in html, "l'utilisateur doit savoir que c'est du texte extrait"


def test_markdown_preview_is_rendered_not_escaped_wholesale(documents):
    html = build_document_preview(str(documents["md"]))
    assert "<h2>" in html or "<h1>" in html
    assert "<strong>" in html


def test_markdown_table_cells_are_rendered_in_preview():
    blocks = parse_markdown("| A | B |\n|---|---|\n| **Taille** | *petite* |")
    html = render_blocks_html(blocks)
    assert "<strong>Taille</strong>" in html
    assert "<em>petite</em>" in html
    assert "**" not in html


def test_file_tables_keep_literal_asterisks(tmp_path):
    """Un « * » d'une vraie cellule de document n'est pas du Markdown."""
    import docx

    path = tmp_path / "calc.docx"
    document = docx.Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Formule"
    table.cell(1, 0).text = "5 * 3 * 2"
    document.save(str(path))
    assert "5 * 3 * 2" in build_document_preview(str(path))


# ── Échappement ────────────────────────────────────────────────────────────


def test_dangerous_content_is_escaped(tmp_path):
    path = tmp_path / "xss.md"
    path.write_text("# <script>alert(1)</script>\n\nTexte & <b>brut</b>.", encoding="utf-8")
    html = build_document_preview(str(path))
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html
    assert "&amp;" in html


def test_filename_is_escaped(tmp_path):
    # « < » et « > » sont interdits dans un nom de fichier Windows : « & »
    # suffit à vérifier que le nom passe bien par html.escape().
    path = tmp_path / "R&D 2024.txt"
    path.write_text("contenu", encoding="utf-8")
    html = build_document_preview(str(path))
    assert "R&amp;D 2024.txt" in html
    assert "R&D 2024.txt" not in html


# ── Robustesse ─────────────────────────────────────────────────────────────


def test_missing_file_gives_a_page_not_an_exception():
    html = build_document_preview("nulle_part/absent.docx")
    assert html.lstrip().startswith("<!DOCTYPE html>")
    assert "introuvable" in html


def test_corrupt_file_gives_a_page_not_an_exception(tmp_path):
    path = tmp_path / "casse.docx"
    path.write_bytes(b"pas un docx")
    html = build_document_preview(str(path))
    assert html.lstrip().startswith("<!DOCTYPE html>")
    assert "sheet" in html


def test_empty_blocks_render_a_message():
    assert "Document vide" in render_blocks_html([])


# ── Résumé de document ─────────────────────────────────────────────────────


def test_document_summary(documents):
    summary = document_summary(str(documents["docx"]))
    assert summary["name"] == "doc.docx"
    assert summary["title"] == "doc"
    assert summary["format"] == "docx"
    assert summary["label"] == "Document Word"
    assert summary["size"] > 0
    assert summary["native_viewer"] is False

    assert document_summary(str(documents["pdf"]))["native_viewer"] is True


# ── Artifacts « document » ─────────────────────────────────────────────────


def test_artifact_from_document(documents):
    artifact = artifact_from_document(str(documents["docx"]))
    assert artifact is not None
    assert artifact.kind == "document"
    assert artifact.is_file is True
    assert artifact.title == "doc.docx"
    assert artifact.language == "docx"


def test_artifact_from_document_refuses_unknown_and_missing(tmp_path):
    assert artifact_from_document("nulle_part/absent.docx") is None
    archive = tmp_path / "a.zip"
    archive.write_bytes(b"PK")
    assert artifact_from_document(str(archive)) is None


def test_artifacts_from_documents_deduplicates(documents):
    path = str(documents["docx"])
    artifacts = artifacts_from_documents([path, path, "absent.docx", None, ""])
    assert len(artifacts) == 1
    assert artifacts[0].index == 0


def test_artifacts_from_documents_continues_indexing(documents):
    artifacts = artifacts_from_documents(
        [str(documents["docx"]), str(documents["pptx"])], start_index=2
    )
    assert [a.index for a in artifacts] == [2, 3]


def test_build_preview_document_dispatches_to_the_document_renderer(documents):
    artifact = artifact_from_document(str(documents["pptx"]))
    html = build_preview_document(artifact)
    assert '<div class="slide">' in html


def test_write_artifact_html_serves_pdf_natively(documents):
    """Edge affiche le PDF lui-même : inutile de le convertir en HTML."""
    served = write_artifact_html(artifact_from_document(str(documents["pdf"])))
    assert served == documents["pdf"]

    served_docx = write_artifact_html(artifact_from_document(str(documents["docx"])))
    assert served_docx.suffix == ".html"
    assert served_docx.exists()
    served_docx.unlink()


def test_html_artifact_detection_is_unchanged():
    """La détection des blocs HTML/SVG ne doit pas avoir bougé."""
    artifacts = detect_artifacts("Voici :\n```html\n<div>Bonjour</div>\n```")
    assert len(artifacts) == 1
    assert artifacts[0].kind == "html"
    assert artifacts[0].is_file is False
    assert detect_artifacts("```python\nprint(1)\n```") == []
