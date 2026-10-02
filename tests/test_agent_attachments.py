"""
Tests de lecture des pièces jointes des agents (processors/attachments.py).

Les deux lecteurs — page Agents du GUI et Relay mobile — appelaient
`process_file`, une méthode qui n'existe ni sur PDFProcessor ni sur
DOCXProcessor : l'agent recevait un message d'erreur (GUI) ou les octets
bruts du fichier (Relay), et un .xlsx joint depuis le GUI était lu comme du
texte. On lit ici de vrais fichiers construits à la volée, par la fonction
commune puis par chacun des deux lecteurs.
"""

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from processors.attachments import read_attachment_text
from relay import agent_relay

# Texte reconnaissable placé dans chaque fichier de test
EXPECTED = {
    "pdf": "Chiffre d'affaires du trimestre",
    "docx": "Compte rendu de réunion",
    "xlsx": "Résine époxy",
    "pptx": "Plan de lancement",
    "csv": "Vis inox",
    "md": "# Notes de version",
}


@pytest.fixture(name="documents")
def _documents(tmp_path):
    """Un fichier par format, chacun contenant son texte de EXPECTED."""
    paths = {kind: tmp_path / f"piece.{kind}" for kind in EXPECTED}

    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), EXPECTED["pdf"])
    pdf.save(str(paths["pdf"]))
    pdf.close()

    word = Document()
    word.add_paragraph(EXPECTED["docx"])
    word.save(str(paths["docx"]))

    workbook = Workbook()
    workbook.active.append(["Produit", "Quantité"])
    workbook.active.append([EXPECTED["xlsx"], 42])
    workbook.save(str(paths["xlsx"]))

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = EXPECTED["pptx"]
    presentation.save(str(paths["pptx"]))

    paths["csv"].write_text(f"article,stock\n{EXPECTED['csv']},12\n", encoding="utf-8")
    paths["md"].write_text(f"{EXPECTED['md']}\n\n- correctif\n", encoding="utf-8")
    return paths


# ── Fonction commune ──────────────────────────────────────────────────────


@pytest.mark.parametrize("kind", EXPECTED)
def test_extracts_text_not_raw_bytes(documents, kind):
    text = read_attachment_text(str(documents[kind]))
    assert EXPECTED[kind] in text
    assert not text.startswith(("%PDF", "PK"))


def test_extension_wins_over_generic_gui_type(documents):
    # Bouton « Code » + filtre « Tous (*.*) » : l'extension choisit le processeur
    assert EXPECTED["pdf"] in read_attachment_text(str(documents["pdf"]), "Code")


@pytest.mark.parametrize("name", ["casse.pdf", "casse.docx", "casse.xlsx", "casse.pptx"])
def test_unreadable_document_raises_instead_of_returning_bytes(tmp_path, name):
    path = tmp_path / name
    path.write_bytes(b"PK\x03\x04 contenu illisible \x00\xff")
    with pytest.raises(ValueError):
        read_attachment_text(str(path))


def test_pdf_without_text_raises(tmp_path):
    path = tmp_path / "vide.pdf"
    pdf = pymupdf.open()
    pdf.new_page()
    pdf.save(str(path))
    pdf.close()
    with pytest.raises(ValueError, match="Aucun texte"):
        read_attachment_text(str(path))


# ── Relay mobile ──────────────────────────────────────────────────────────


def test_relay_task_contains_document_text(documents):
    """Tâche envoyée à un agent lancé depuis le mobile avec deux pièces jointes."""
    relay = agent_relay.AgentRelayService.__new__(agent_relay.AgentRelayService)
    task = relay._augment_task_with_files(
        "Résume ces fichiers.", None, [str(documents["pdf"]), str(documents["docx"])]
    )
    assert EXPECTED["pdf"] in task
    assert EXPECTED["docx"] in task
    assert "%PDF" not in task and "PK\x03\x04" not in task


def test_relay_reports_unreadable_document(tmp_path, monkeypatch):
    monkeypatch.setattr(agent_relay.logger, "disabled", True)  # rien dans logs/
    path = tmp_path / "casse.pdf"
    path.write_bytes(b"%PDF-1.7\n1 0 obj\n<</Type/Catalog")
    text = agent_relay.AgentRelayService._read_attached_file(str(path))
    assert text.startswith("[Erreur de lecture :")
    assert "Type/Catalog" not in text


# ── Page Agents du GUI ────────────────────────────────────────────────────


def test_gui_reader_reads_each_document_type(documents):
    """Type du bouton choisi dans la page Agents + fichier sélectionné."""
    pytest.importorskip("tkinter")
    from interfaces.agents.file_handling import FileHandlingMixin

    read = FileHandlingMixin._read_attached_file
    assert EXPECTED["pdf"] in read(str(documents["pdf"]), "PDF")
    assert EXPECTED["docx"] in read(str(documents["docx"]), "DOCX")
    assert EXPECTED["xlsx"] in read(str(documents["xlsx"]), "Excel")
    assert EXPECTED["pptx"] in read(str(documents["pptx"]), "PowerPoint")
